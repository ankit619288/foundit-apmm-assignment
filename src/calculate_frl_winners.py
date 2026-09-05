from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parents[1]
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "outputs"
NOTES_DIR = BASE_DIR / "notes"

PURCHASE_FILE = DATA_DIR / "purchase.xlsx"
OUTPUT_FILE = OUTPUT_DIR / "fRL_winners.xlsx"
METHOD_NOTE_FILE = NOTES_DIR / "frl_method_note.md"

REFERENCE_DATE = datetime(2026, 2, 20)
TOP_N_CATEGORY = 30
TOP_N_MVP = 5

OUTPUT_COLUMNS = [
    "Source",
    "Category",
    "User Type",
    "MVP Rank",
    "MVP Score",
    "Company Name",
    "Company Type",
    "Service Channel",
    "Login",
    "Email",
    "PC",
    "OC",
    "JP",
    "Start Date",
    "End Date",
]

REQUIRED_COLUMNS = {
    "company_name",
    "company_type",
    "account_type",
    "service_channel",
    "login",
    "email",
    "PC",
    "OC",
    "JP",
    "pc_start_date",
    "pc_end_date",
}


def classify_user(start_date):
    if pd.isna(start_date):
        return "Existing"

    start_date = pd.to_datetime(start_date, errors="coerce")

    if pd.isna(start_date):
        return "Existing"

    if datetime(2025, 10, 1) <= start_date <= datetime(2026, 1, 31):
        return "New"

    return "Existing"


def load_purchase_data():
    if not PURCHASE_FILE.exists():
        raise FileNotFoundError(f"Purchase input not found: {PURCHASE_FILE}")

    df = pd.read_excel(PURCHASE_FILE)

    missing_columns = sorted(REQUIRED_COLUMNS - set(df.columns))
    if missing_columns:
        raise ValueError(
            "purchase.xlsx is missing required columns: "
            + ", ".join(missing_columns)
        )

    df["start_date_internal"] = pd.to_datetime(df["pc_start_date"], errors="coerce")
    df["end_date_internal"] = pd.to_datetime(df["pc_end_date"], errors="coerce")

    text_columns = [
        "company_name",
        "company_type",
        "account_type",
        "service_channel",
        "login",
        "email",
    ]

    for col in text_columns:
        df[col] = df[col].fillna("").astype(str).str.strip()

    for col in ["PC", "OC", "JP"]:
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)

    return df


def apply_eligibility_rules(df):
    original_count = len(df)
    mask = pd.Series(True, index=df.index)

    mask &= ~df["account_type"].str.lower().isin(["free_trial", "test_job"])

    mask &= ~df["login"].str.lower().str.contains(
        "scrape|_jobs|ftp",
        na=False,
        regex=True,
    )

    mask &= ~df["service_channel"].str.lower().str.contains(
        "scrape",
        na=False,
        regex=True,
    )

    mask &= ~df["company_name"].str.lower().str.contains(
        "immigration|ankit sharma proprietor|freelancer",
        na=False,
        regex=True,
    )

    # Match the supplied winner.py behavior: an account must have a valid end
    # date on or after the reference date to be treated as active.
    mask &= df["end_date_internal"].notna()
    mask &= df["end_date_internal"] >= REFERENCE_DATE

    eligible_df = df[mask].copy()
    eligible_df["User Type"] = eligible_df["start_date_internal"].apply(classify_user)

    print(f"Original rows: {original_count}")
    print(f"Eligible rows: {len(eligible_df)}")
    print(f"Removed rows: {original_count - len(eligible_df)}")

    return eligible_df


def build_output_row(row, category, mvp_rank="", mvp_score=""):
    return {
        "Source": "Purchase",
        "Category": category,
        "User Type": row["User Type"],
        "MVP Rank": mvp_rank,
        "MVP Score": mvp_score,
        "Company Name": row["company_name"],
        "Company Type": row["company_type"],
        "Service Channel": row["service_channel"],
        "Login": row["login"],
        "Email": row["email"],
        "PC": row["PC"],
        "OC": row["OC"],
        "JP": row["JP"],
        "Start Date": row["start_date_internal"],
        "End Date": row["end_date_internal"],
    }


def get_category_winners(df):
    category_map = {
        "Search Smasher": "PC",
        "Campaign Captain": "OC",
        "Posting Champion": "JP",
    }

    winners = []

    for category, metric in category_map.items():
        top_30 = df.sort_values(
            metric,
            ascending=False,
            kind="mergesort",
        ).head(TOP_N_CATEGORY)

        for user_type in ["New", "Existing"]:
            subset = top_30[top_30["User Type"] == user_type]

            if not subset.empty:
                winner = subset.iloc[0]
                winners.append(build_output_row(winner, category))

    return winners


def get_mvp_winners(df):
    mvp_df = df[
        (df["PC"] > 0) &
        (df["OC"] > 0) &
        (df["JP"] > 0)
    ].copy()

    mvp_df["MVP Score"] = mvp_df["PC"] + mvp_df["OC"] + mvp_df["JP"]
    mvp_df = mvp_df.sort_values(
        "MVP Score",
        ascending=False,
        kind="mergesort",
    )

    winners = []

    for user_type in ["New", "Existing"]:
        subset = mvp_df[mvp_df["User Type"] == user_type].copy()
        subset = subset.head(TOP_N_MVP)

        for rank, (_, row) in enumerate(subset.iterrows(), start=1):
            winners.append(
                build_output_row(
                    row,
                    category="MVP",
                    mvp_rank=rank,
                    mvp_score=row["MVP Score"],
                )
            )

    return winners


def describe_category_cutoff_ties(df):
    tie_details = []
    for metric in ["PC", "OC", "JP"]:
        ranked = df.sort_values(metric, ascending=False, kind="mergesort")
        if len(ranked) < TOP_N_CATEGORY:
            continue

        cutoff = ranked.iloc[TOP_N_CATEGORY - 1][metric]
        tied_accounts = int((ranked[metric] == cutoff).sum())
        if tied_accounts > 1:
            tie_details.append(
                f"{metric}'s rank-{TOP_N_CATEGORY} cutoff is {cutoff:g} with "
                f"{tied_accounts} tied eligible accounts"
            )

    if not tie_details:
        return "No category top-30 cutoff ties were found"
    return "; ".join(tie_details)


def save_method_note(cleaning_summary, cutoff_tie_note):
    note = f"""# fRL Winner Calculation Method Note

- Method: Python standardized text, dates, and PC/OC/JP values, then applied every supplied Purchase eligibility rule before ranking.
- Cleaning: {cleaning_summary["original_rows"]:,} source rows became {cleaning_summary["eligible_rows"]:,} eligible rows; {cleaning_summary["removed_rows"]:,} rows were excluded.
- Classification: New means a start date from 1 October 2025 through 31 January 2026; all other or blank start dates are Existing.
- Winners: each PC/OC/JP category uses the overall top 30 and selects the highest New and Existing account; MVP requires all three metrics above zero and takes the top five per user type by PC + OC + JP.
- Anomalies/assumptions: {cleaning_summary["missing_start_dates"]:,} start dates and {cleaning_summary["missing_end_dates"]:,} end dates are blank/invalid; blank starts count as Existing, while blank ends are ineligible because active status cannot be verified, matching the supplied `winner.py`. No old-winners list was provided. {cutoff_tie_note}; stable source order resolves tied values because no secondary business tie-break was supplied.
"""
    METHOD_NOTE_FILE.write_text(note, encoding="utf-8")


def polish_excel_file():
    wb = load_workbook(OUTPUT_FILE)
    ws = wb.active
    ws.title = "fRL Winners"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    wrapped_left = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 24

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column == 6:
                cell.alignment = wrapped_left
            elif cell.column in [4, 5, 11, 12, 13]:
                cell.alignment = center
            else:
                cell.alignment = left

            if cell.value is None:
                cell.value = ""

            if cell.column in [14, 15] and cell.value:
                cell.number_format = "dd-mmm-yyyy"

    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_length = 0
        column_letter = get_column_letter(col_idx)

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))

        adjusted_width = min(max(max_length + 2, 12), 38)
        ws.column_dimensions[column_letter].width = adjusted_width

    ws.column_dimensions["F"].width = 48
    ws.column_dimensions["J"].width = 34
    ws.column_dimensions["N"].width = 16
    ws.column_dimensions["O"].width = 16

    for row_idx in range(2, ws.max_row + 1):
        company_name = str(ws.cell(row=row_idx, column=6).value or "")
        if len(company_name) > 48:
            ws.row_dimensions[row_idx].height = 32

    wb.save(OUTPUT_FILE)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    NOTES_DIR.mkdir(parents=True, exist_ok=True)

    df = load_purchase_data()
    original_rows = len(df)

    eligible_df = apply_eligibility_rules(df)

    cleaning_summary = {
        "original_rows": original_rows,
        "eligible_rows": len(eligible_df),
        "removed_rows": original_rows - len(eligible_df),
        "missing_start_dates": int(df["start_date_internal"].isna().sum()),
        "missing_end_dates": int(df["end_date_internal"].isna().sum()),
    }

    category_winners = get_category_winners(eligible_df)
    mvp_winners = get_mvp_winners(eligible_df)

    final_df = pd.DataFrame(
        category_winners + mvp_winners,
        columns=OUTPUT_COLUMNS,
    )

    final_df["MVP Rank"] = final_df["MVP Rank"].replace("", "-")
    final_df["MVP Score"] = final_df["MVP Score"].replace("", "-")

    final_df.to_excel(OUTPUT_FILE, index=False)
    polish_excel_file()
    save_method_note(cleaning_summary, describe_category_cutoff_ties(eligible_df))

    print("\nfRL winners generated successfully.")
    print(f"Output Excel: {OUTPUT_FILE}")
    print(f"Method note: {METHOD_NOTE_FILE}")
    print("\nWinner preview:")
    print(final_df.to_string(index=False))


if __name__ == "__main__":
    main()
